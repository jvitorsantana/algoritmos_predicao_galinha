"""
Exporta as metricas completas do Experimento 3 para um arquivo Excel.

Le os JSONs gerados por comparacao_sexo.py e comparacao_peso.py:
  - results/metricas_completas_sexo.json  (classificacao)
  - results/metricas_completas_peso.json  (regressao)

Gera results/metricas_experimento_3.xlsx com abas:
  - Resumo                : visao geral dos dois tipos de tarefa
  - Classificacao_Metricas: Accuracy/Precision/Recall/F1 por modelo (sexo)
  - Classificacao_MatrizConfusao: matriz de confusao de cada modelo
  - Regressao_Metricas    : R2/RMSE/MAE por modelo (peso)
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent.parent
RESULTS = ROOT / 'results'

HEADER_FILL = PatternFill('solid', fgColor='4472C4')
SUB_FILL = PatternFill('solid', fgColor='D9E1F2')
DIAG_FILL = PatternFill('solid', fgColor='C6EFCE')
HEADER_FONT = Font(bold=True, color='FFFFFF')
BOLD = Font(bold=True)
CENTER = Alignment(horizontal='center', vertical='center')
THIN = Side(style='thin', color='BFBFBF')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_header(ws, row, n_cols, start_col=1):
    for c in range(start_col, start_col + n_cols):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER


def autosize(ws, min_width=10, max_width=42):
    for col in ws.columns:
        length = max((len(str(c.value)) for c in col if c.value is not None), default=0)
        letter = get_column_letter(col[0].column)
        ws.column_dimensions[letter].width = max(min_width, min(max_width, length + 2))


def fmt(v, nd=4):
    return round(v, nd) if isinstance(v, (int, float)) else v


sexo = json.loads((RESULTS / 'metricas_completas_sexo.json').read_text(encoding='utf-8'))
peso = json.loads((RESULTS / 'metricas_completas_peso.json').read_text(encoding='utf-8'))

wb = Workbook()

# ---------------------------------------------------------------------------
# Aba: Resumo
# ---------------------------------------------------------------------------
ws = wb.active
ws.title = 'Resumo'
ws['A1'] = 'Experimento 3 - Metricas Completas dos Modelos'
ws['A1'].font = Font(bold=True, size=14)
ws['A3'] = ('Classificacao (SEXO): possui Matriz de Confusao, Accuracy, Precision, '
            'Recall e F1-score.')
ws['A4'] = ('Regressao (PESO): NAO possui matriz de confusao nem '
            'Accuracy/Precision/Recall/F1. As metricas equivalentes sao R2, RMSE e MAE.')
ws['A6'] = 'Conteudo das abas:'
ws['A6'].font = BOLD
abas = [
    ('Classificacao_Metricas', f"{len(sexo['models'])} classificadores - alvo {sexo['target']} "
     f"(classes: {', '.join(sexo['classes'])})"),
    ('Classificacao_MatrizConfusao', 'Matriz de confusao de cada classificador (linha=real, coluna=previsto)'),
    ('Regressao_Metricas', f"{len(peso['models'])} regressores - alvo {peso['target']}"),
]
r = 7
for nome, desc in abas:
    ws.cell(row=r, column=1, value=nome).font = BOLD
    ws.cell(row=r, column=2, value=desc)
    r += 1
ws.cell(row=r + 1, column=1,
        value=f"Baseline classificacao (classe majoritaria): Acc = {fmt(sexo['baseline_acc'])}")
ws.cell(row=r + 2, column=1, value=f"N amostras de teste - classificacao: {sexo['n_test']} | regressao: {peso['n_test']}")
autosize(ws)

# ---------------------------------------------------------------------------
# Aba: Classificacao_Metricas
# ---------------------------------------------------------------------------
ws = wb.create_sheet('Classificacao_Metricas')
classes = sexo['classes']
# Metricas GERAIS (consideram as duas classes): Precision/Recall/F1 = media
# ponderada (weighted) pelo numero de amostras de cada classe; pareia com a Accuracy.
ws['A1'] = ('Metricas gerais por modelo. Precision/Recall/F1 = media ponderada '
            '(weighted) das duas classes (Femea e Macho).')
ws['A1'].font = Font(italic=True, size=10)
headers = ['Modelo', 'Tipo', 'Accuracy', 'Precision', 'Recall', 'F1-score']
hdr_row = 3
for j, h in enumerate(headers, start=1):
    ws.cell(row=hdr_row, column=j, value=h)
style_header(ws, hdr_row, len(headers))
# ordena por F1 geral (weighted) desc
for m in sorted(sexo['models'], key=lambda x: x['f1_weighted'], reverse=True):
    ws.append([
        m['name'], 'Classificacao', fmt(m['accuracy']),
        fmt(m['precision_weighted']), fmt(m['recall_weighted']), fmt(m['f1_weighted']),
    ])
for row in ws.iter_rows(min_row=hdr_row + 1, max_row=ws.max_row):
    for cell in row:
        cell.border = BORDER
        if cell.column > 2:
            cell.alignment = CENTER
# linha de baseline
br = ws.max_row + 2
ws.cell(row=br, column=1, value='Baseline (classe majoritaria)').font = BOLD
ws.cell(row=br, column=3, value=fmt(sexo['baseline_acc'])).alignment = CENTER
autosize(ws)

# ---------------------------------------------------------------------------
# Aba: Classificacao_MatrizConfusao
# ---------------------------------------------------------------------------
ws = wb.create_sheet('Classificacao_MatrizConfusao')
ws['A1'] = 'Matrizes de Confusao por Modelo (linha = REAL, coluna = PREVISTO)'
ws['A1'].font = Font(bold=True, size=12)
row = 3
for m in sorted(sexo['models'], key=lambda x: x['f1_pos'], reverse=True):
    ws.cell(row=row, column=1, value=m['name']).font = Font(bold=True, size=11)
    row += 1
    # cabecalho de colunas previstas
    ws.cell(row=row, column=2, value='Previsto ->').font = BOLD
    for j, cls in enumerate(classes):
        c = ws.cell(row=row, column=3 + j, value=cls)
        c.fill = SUB_FILL
        c.font = BOLD
        c.alignment = CENTER
        c.border = BORDER
    ws.cell(row=row, column=3 + len(classes), value='Total real').font = BOLD
    row += 1
    cm = m['confusion_matrix']
    for i, cls in enumerate(classes):
        rc = ws.cell(row=row, column=2, value=f'Real: {cls}')
        rc.fill = SUB_FILL
        rc.font = BOLD
        rc.border = BORDER
        for j in range(len(classes)):
            c = ws.cell(row=row, column=3 + j, value=cm[i][j])
            c.alignment = CENTER
            c.border = BORDER
            if i == j:
                c.fill = DIAG_FILL  # acertos na diagonal
        ws.cell(row=row, column=3 + len(classes), value=sum(cm[i])).alignment = CENTER
        row += 1
    # totais previstos
    tc = ws.cell(row=row, column=2, value='Total previsto')
    tc.font = BOLD
    for j in range(len(classes)):
        col_total = sum(cm[i][j] for i in range(len(classes)))
        ws.cell(row=row, column=3 + j, value=col_total).alignment = CENTER
    ws.cell(row=row, column=3 + len(classes),
            value=sum(sum(r) for r in cm)).alignment = CENTER
    row += 2
autosize(ws)

# ---------------------------------------------------------------------------
# Aba: Regressao_Metricas
# ---------------------------------------------------------------------------
ws = wb.create_sheet('Regressao_Metricas')
ws['A1'] = 'Regressao (PESO) - sem matriz de confusao / Accuracy / Precision / Recall / F1'
ws['A1'].font = Font(bold=True, size=11)
ws['A2'] = 'Metricas equivalentes para regressao: R2, RMSE (g) e MAE (g).'
headers = ['Modelo', 'Tipo', 'R2 (CV)', 'R2 (Teste)', 'RMSE (g)', 'MAE (g)']
ws.append([])
ws.append(headers)
hdr_row = ws.max_row
style_header(ws, hdr_row, len(headers))
for m in sorted(peso['models'], key=lambda x: x['r2_test'], reverse=True):
    ws.append([
        m['name'], 'Regressao', fmt(m['r2_cv']), fmt(m['r2_test']),
        fmt(m['rmse'], 2), fmt(m['mae'], 2),
    ])
for row in ws.iter_rows(min_row=hdr_row + 1, max_row=ws.max_row):
    for cell in row:
        cell.border = BORDER
        if cell.column > 2:
            cell.alignment = CENTER
autosize(ws)

OUT = RESULTS / 'metricas_experimento_3.xlsx'
wb.save(OUT)
print(f'Excel salvo em: {OUT}')
